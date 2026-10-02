from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from upgrade_algorithm_report_v11_1 import add_paragraph_before, move_before
from upgrade_algorithm_report_v11_3 import align_table_indents
from upgrade_algorithm_report_v11_4 import add_table_before


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "为什么数学建模能够支撑供热转型决策_算法增强版V7.docx"
OUTPUT = ROOT / "为什么数学建模能够支撑供热转型决策_算法增强版V8.docx"
FIGURE = ROOT / "outputs" / "v11_6_figures" / "multiyear_robustness_dashboard.png"


def refresh_summary(doc) -> None:
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            if "为什么可信：V11.5" in paragraph.text and "V11.5" in run.text:
                run.text = run.text.replace("V11.5", "V11.6")
            if "全工程68项" in paragraph.text and "68" in run.text:
                run.text = run.text.replace("68", "73")


def add_picture_before(doc, anchor, path: Path) -> None:
    picture = doc.add_paragraph()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.paragraph_format.keep_with_next = True
    shape = picture.add_run().add_picture(str(path), width=Inches(6.35))
    shape._inline.docPr.set(
        "descr",
        "2020至2024五年联合代表日覆盖告警、跨年容量通过率、七个极端年求解时间以及成本排放分布",
    )
    shape._inline.docPr.set("title", "V11.6多气象年鲁棒性证据")
    move_before(picture, anchor)
    caption = add_paragraph_before(
        doc,
        anchor,
        "图12  五年联合代表日、跨年容量审计与极端年正式回放结果",
        "Caption" if "Caption" in doc.styles else "Normal",
        WD_ALIGN_PARAGRAPH.CENTER,
    )
    caption.paragraph_format.space_after = Pt(8)


def add_v11_6_section(doc) -> None:
    anchor = next(
        paragraph
        for paragraph in doc.paragraphs
        if paragraph.text.strip().startswith("3  为什么该模型具有工程物理可信性")
    )
    add_paragraph_before(
        doc,
        anchor,
        "2.7.12  V11.6：五年联合代表日与同一投资方案跨年审计",
        "Heading 3",
    )
    add_paragraph_before(
        doc,
        anchor,
        "单一气象年可能低估年际热峰、冷峰、爬坡、最低温和热泵最差工况。V11.6把乌鲁木齐、固原和武汉2020—2024年NASA POWER逐时气象共同纳入代表日生成：五年共1827个完整日先聚类，再强制保留逐负荷峰值与爬坡、最低温、最低资源可用率、最大有效负荷和最差热泵电热比日；权重归一到平均年365.4天，并按五年平均热冷需求重新校准。三场景最终只需13、12和17个代表日，跨五年覆盖告警由10、8、4项全部降为0，峰值与爬坡捕获率均为100%。",
        "Normal",
    )
    add_picture_before(doc, anchor, FIGURE)
    add_paragraph_before(
        doc,
        anchor,
        "代表模型只求解一次投资，随后把热源、管网、储热、配电和微网方案全部冻结，逐年执行可靠容量审计。原2023单年方案的五年通过数为S1 5/5、S2 1/5、S3 4/5；五年联合方案提升为三场景全部5/5。场景二的配电可靠容量采用组合调度口径：在15%热备用率下选择可行峰值热源组合，再计算其最小可靠电力需求，避免把互为备用的热泵和电锅炉重复假设为同时满发。",
        "Normal",
    )
    add_table_before(
        doc,
        anchor,
        ["场景/年份", "小时", "变量", "整数变量", "求解/s", "热/冷失供", "证书"],
        [
            ("S1 乌鲁木齐 2024", "8784", "440,615", "8,787", "24.28", "0 / 0", "通过"),
            ("S2 固原 2021", "8760", "424,490", "26,280", "425.16", "0 / 0", "通过"),
            ("S2 固原 2024", "8784", "423,050", "26,352", "26.03", "0 / 0", "通过"),
            ("S3 武汉 2021", "8760", "314,883", "26,280", "62.14", "0 / 0", "通过"),
            ("S3 武汉 2022", "8760", "314,099", "26,280", "53.83", "0 / 0", "通过"),
            ("S3 武汉 2023", "8760", "313,579", "26,280", "44.86", "0 / 0", "通过"),
            ("S3 武汉 2024", "8784", "314,011", "26,352", "49.08", "0 / 0", "通过"),
        ],
        [2100, 900, 1350, 1350, 1000, 1360, 1300],
    )
    add_paragraph_before(
        doc,
        anchor,
        "七个由最低温、最大热峰、最大爬坡、最大冷峰和可靠电力需求筛选出的极端年，均完成连续8760/8784小时正式回放并取得机器可读验收证书。所有年份热失供率为0，武汉冷失供率亦为0；同一场景的多个年份复用完全相同的冻结投资方案，而不是逐年重新扩容。",
        "Normal",
    )

    add_paragraph_before(
        doc,
        anchor,
        "2.7.13  CHP源侧弃热、储热退化成本与同步充放验收",
        "Heading 3",
    )
    add_paragraph_before(
        doc,
        anchor,
        "场景一具有CHP最小电出力约束。原模型没有显式冷凝/旁路弃热通道；如果仅为提速而放松储热充放互斥，优化器会利用往返损耗循环消纳多余热量。V11.6没有接受该低成本解，而是新增三层物理约束：第一，只有CHP源侧可以使用受限弃热，弃热计入机组容量、热电耦合、燃料成本、排放和㶲损失，但不计入有效供热或清洁热份额；第二，储热充放量计入泵耗、换热器循环和寿命退化代理成本；第三，正式证书逐时计算min(充热功率, 放热功率)，只要最大同步充放超过数值容差就拒绝认证。",
        "Normal",
    )
    add_paragraph_before(
        doc,
        anchor,
        "修复后的S1 2024回放求解24.28秒，中央水罐最大同步充放功率为0、状态方程残差为0、循环边界通过；年弃热约800.4 MWh。该数值不是项目实测结论，而是需要在现场核实的工程量：正式可研必须确认冷凝/旁路能力、弃热环境约束、最小发电曲线和储热O&M成本。如果实测不支持当前180 kW弃热上限，应重新规划，而不能用储热循环损耗替代物理设备。",
        "Normal",
    )
    add_paragraph_before(
        doc,
        anchor,
        "V11.6全工程73项自动化测试通过。新增测试覆盖普通年与闰年、固定方案复用、跨年权重守恒与极端捕获、CHP最小发电下的源侧弃热，以及储热同步充放拒绝。该版本把结论从‘三场景单年可行’推进为‘三场景五年容量审计全通过、七个极端年完整时序认证全通过’，并保留公开代理数据不能替代实测的适用边界。",
        "Normal",
    )


def main() -> None:
    for path in (SOURCE, FIGURE):
        if not path.exists():
            raise FileNotFoundError(path)
    doc = Document(SOURCE)
    refresh_summary(doc)
    add_v11_6_section(doc)
    align_table_indents(doc)
    doc.core_properties.title = "为什么数学建模能够支撑供热转型决策——算法增强版V8"
    doc.core_properties.subject = "V11.6多气象年鲁棒规划、极端年回放与储热物理验收"
    doc.core_properties.version = "11.6"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
