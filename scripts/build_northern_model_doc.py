from __future__ import annotations

import json
import re
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "北方城镇供热源荷匹配与规划寻优模型说明书.md"
OUTPUT = ROOT / "docs" / "北方城镇供热源荷匹配与规划寻优模型说明书.docx"
FIG = ROOT / "outputs" / "figures"

BLUE = "17324D"
MID_BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
GRAY = "5B6573"
LIGHT = "F4F6F9"
PALE_BLUE = "E8EEF5"
GOLD = "B7791F"
WHITE = "FFFFFF"
INK = "20262E"
TABLE_WIDTH = 9360
TABLE_INDENT = 120


def shade(element, fill: str) -> None:
    props = element.get_or_add_tcPr() if element.tag.endswith("tc") else element.get_or_add_pPr()
    shd = props.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        props.append(shd)
    shd.set(qn("w:fill"), fill)


def set_run_font(run, latin="Calibri", east_asia="Noto Sans SC", size=None, color=None, bold=None, italic=None):
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(sum(widths)))
    tblW.set(qn("w:type"), "dxa")
    tblInd = tblPr.find(qn("w:tblInd"))
    if tblInd is None:
        tblInd = OxmlElement("w:tblInd")
        tblPr.append(tblInd)
    tblInd.set(qn("w:w"), str(TABLE_INDENT))
    tblInd.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        trPr.append(cant_split)
        for i, cell in enumerate(row.cells):
            tcW = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tcW)
            tcW.set(qn("w:w"), str(widths[i]))
            tcW.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    header_pr = table.rows[0]._tr.get_or_add_trPr()
    table_header = OxmlElement("w:tblHeader")
    table_header.set(qn("w:val"), "true")
    header_pr.append(table_header)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=GRAY)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)
    run = paragraph.add_run(" 页")
    set_run_font(run, size=9, color=GRAY)


def configure_styles(doc: Document):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans SC")
    pf = normal.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.space_before = Pt(0)
    pf.space_after = Pt(8)
    pf.line_spacing = 1.333

    for name, size, color, before, after in [
        ("Heading 1", 16, MID_BLUE, 18, 10),
        ("Heading 2", 13, MID_BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans SC")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(10.5)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans SC")
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.194)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.208

    if "Equation" not in styles:
        eq = styles.add_style("Equation", WD_STYLE_TYPE.PARAGRAPH)
    else:
        eq = styles["Equation"]
    eq.font.name = "Cambria Math"
    eq.font.size = Pt(10.5)
    eq.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    eq.paragraph_format.space_before = Pt(4)
    eq.paragraph_format.space_after = Pt(8)
    eq.paragraph_format.keep_together = True

    if "Lead Callout" not in styles:
        callout = styles.add_style("Lead Callout", WD_STYLE_TYPE.PARAGRAPH)
    else:
        callout = styles["Lead Callout"]
    callout.font.name = "Calibri"
    callout.font.size = Pt(11)
    callout.font.bold = True
    callout.font.color.rgb = RGBColor.from_string(BLUE)
    callout._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans SC")
    callout.paragraph_format.left_indent = Inches(0.18)
    callout.paragraph_format.right_indent = Inches(0.18)
    callout.paragraph_format.space_before = Pt(6)
    callout.paragraph_format.space_after = Pt(10)
    callout.paragraph_format.line_spacing = 1.2


def add_inline_markup(paragraph, text: str, size=10.5):
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        content = part[2:-2] if bold else part
        run = paragraph.add_run(content)
        set_run_font(run, size=size, bold=bold)


def add_caption(doc, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    set_run_font(run, size=9, color=GRAY, italic=False)


def add_figure(doc, path: Path, caption: str, width=6.15):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    picture = p.add_run().add_picture(str(path), width=Inches(width))
    picture._inline.docPr.set("descr", caption)
    picture._inline.docPr.set("title", caption)
    add_caption(doc, caption)


def add_result_table(doc):
    rows = [
        ("年化综合成本", "79.51 万元", "包含投资、运行、碳、温位与舒适成本"),
        ("年碳排", "310.58 tCO₂", "示例典型日折算结果"),
        ("供热缺额率", "0", "满足北方民生保供硬约束"),
        ("火用效率", "82.51%", "体现温位匹配效果"),
        ("新增余热回收", "450 kW", "优先消纳工业余热"),
        ("新增余热源热泵", "约 211 kW", "承担常态清洁供热"),
        ("新增低温储热", "约 1235 kWh", "削峰并降低高价电时段用电"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    headers = ["指标", "示例结果", "含义"]
    for i, value in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = value
        shade(cell._tc, PALE_BLUE)
        for run in cell.paragraphs[0].runs:
            set_run_font(run, size=9.5, bold=True, color=BLUE)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for item in rows:
        cells = table.add_row().cells
        for i, value in enumerate(item):
            cells[i].text = value
            for run in cells[i].paragraphs[0].runs:
                set_run_font(run, size=9.2)
            cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if i < 2 else WD_ALIGN_PARAGRAPH.LEFT
    set_table_geometry(table, [2100, 1800, 5460])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_references(doc):
    doc.add_heading("参考文献", level=1)
    refs = [
        "[1] 江亿, 胡珊, 付林, 等. 构建新型零碳热力系统[J]. 中国工程科学, 2025, 27(5).",
        "[2] Lin Q, Lu S, Yue L, et al. Research on low-carbon district heating system planning approach based on thermal energy grade and integrated demand response[J]. Energy and Buildings, 2025, 337: 115660.",
        "[3] Ceruti A, Lambert J, Spliethoff H. Integrating renewable energy and thermal storage in district heating networks: A design optimization approach[J]. Energy Conversion and Management, 2025, 345: 120323.",
        "[4] Chen J, Lv H, Ding L, et al. Multi-objective optimization for the operation of district heating systems considering flexibility and economy[J]. Applied Thermal Engineering, 2025, 263: 125395.",
        "[5] Lyu Z, Huang M, Liu B, et al. Robust optimization of regional integrated energy systems considering thermal comfort and uncertainty in heat networks[J]. Applied Thermal Engineering, 2025, 275: 126895.",
        "[6] 崔梦颖, 殷勇高, 赵栋霖, 等. 地埋管跨季节储热耦合热泵供暖系统性能与经济性分析[J]. 储能科学与技术, 2026, 15(3): 983-992.",
        "[7] 崔杨, 等. 考虑风电消纳的区域综合能源系统源荷协调经济调度[J]. 电网技术, 2020.",
        "[8] 李仲博. 基于数字孪生的城镇供热系统按需精准调控技术研究与应用[D]. 浙江大学, 2023.",
    ]
    for ref in refs:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.first_line_indent = Inches(-0.2)
        p.paragraph_format.space_after = Pt(4)
        add_inline_markup(p, ref, size=9.2)


def add_toc(doc):
    doc.add_heading("目录", level=1)
    entries = [
        ("摘要", 3), ("1 选题定位与北方问题特征", 3), ("2 总体技术路线", 4),
        ("3 基本集合、参数与变量", 5), ("4 目标函数与多目标决策", 6),
        ("5 北方专属物理模型", 6), ("6 节点热平衡与容量约束", 9),
        ("7 不确定性与复合寒潮测试", 10), ("8 算法与求解流程", 11),
        ("9 示例结果与解释", 12), ("10 方案的核心优势", 14),
        ("11 与常见方案的对比", 15), ("12 评价指标体系", 16),
        ("13 数据需求与参数校准", 16), ("14 局限性与后续升级", 16),
        ("15 结论", 16), ("参考文献", 18),
    ]
    for title, page in entries:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.15)
        p.paragraph_format.right_indent = Inches(0.15)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(6.1))
        run = p.add_run(title)
        set_run_font(run, size=10.5, color=INK)
        run = p.add_run("\t" + str(page))
        set_run_font(run, size=10, color=GRAY)


def add_cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(40)
    p.paragraph_format.space_after = Pt(16)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("CP-202611｜模型技术说明书")
    set_run_font(run, size=11, bold=True, color=GOLD)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run("面向北方严寒地区的城镇供热\n源荷匹配与规划寻优模型")
    set_run_font(run, size=27, bold=True, color=BLUE, east_asia="Noto Sans SC")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(28)
    run = p.add_run("温位—时间—空间—风险四维协同｜源—网—荷—储联合优化")
    set_run_font(run, size=13, color=DARK_BLUE, bold=True, east_asia="Noto Sans SC")

    add_figure(doc, FIG / "northern_framework.png", "北方城镇供热四维协同总体框架", width=5.8)

    p = doc.add_paragraph(style="Lead Callout")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("核心主张：常态工况以余热和热泵实现低碳供热，极端寒潮依靠可靠容量和储热确保零失供，并量化保供的经济与碳代价。")
    shade(p._p, LIGHT)

def is_equation(line: str) -> bool:
    return len(line) < 120 and any(token in line for token in (" = ", " ≤ ", " ≥ ", "min ", "max")) and not line.startswith("-")


def new_decimal_numbering(doc: Document) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start"); start.set(qn("w:val"), "1"); lvl.append(start)
    fmt = OxmlElement("w:numFmt"); fmt.set(qn("w:val"), "decimal"); lvl.append(fmt)
    text = OxmlElement("w:lvlText"); text.set(qn("w:val"), "%1."); lvl.append(text)
    jc = OxmlElement("w:lvlJc"); jc.set(qn("w:val"), "left"); lvl.append(jc)
    pPr = OxmlElement("w:pPr")
    ind = OxmlElement("w:ind"); ind.set(qn("w:left"), "540"); ind.set(qn("w:hanging"), "280"); pPr.append(ind)
    lvl.append(pPr); abstract.append(lvl); numbering.append(abstract)
    num = OxmlElement("w:num"); num.set(qn("w:numId"), str(num_id))
    aid = OxmlElement("w:abstractNumId"); aid.set(qn("w:val"), str(abstract_id)); num.append(aid); numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int):
    pPr = paragraph._p.get_or_add_pPr()
    numPr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl"); ilvl.set(qn("w:val"), "0"); numPr.append(ilvl)
    nid = OxmlElement("w:numId"); nid.set(qn("w:val"), str(num_id)); numPr.append(nid)
    pPr.append(numPr)


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(1.0)
    sec.bottom_margin = Inches(1.0)
    sec.left_margin = Inches(1.0)
    sec.right_margin = Inches(1.0)
    sec.header_distance = Inches(0.492)
    sec.footer_distance = Inches(0.492)
    configure_styles(doc)

    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = header.add_run("CP-202611  北方城镇供热源荷匹配与规划寻优")
    set_run_font(run, size=8.5, color=GRAY)
    add_page_number(sec.footer.paragraphs[0])

    add_cover(doc)
    doc.add_page_break()
    add_toc(doc)
    doc.add_page_break()

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    i = 0
    active_num_id = None
    skip_title = True
    figure_after = {
        "2 总体技术路线": (FIG / "northern_framework.png", "图1  北方源—网—荷—储四维协同规划框架"),
        "5.2 空气源热泵低温性能": (FIG / "air_hp_cold_performance.png", "图2  严寒条件下空气源热泵COP与可用容量示意"),
        "7.3 冻结投资压力测试": (FIG / "cold_wave_stress.png", "图3  复合寒潮冻结投资压力测试结果"),
        "9 示例结果与解释": (ROOT / "outputs" / "experiments" / "pareto_front.png", "图4  成本—碳排Pareto前沿与推荐膝点"),
    }
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            active_num_id = None
            i += 1
            continue
        if line.startswith("# "):
            if skip_title:
                skip_title = False
            i += 1
            continue
        if line.startswith("## "):
            title = line[3:].strip()
            doc.add_heading(title, level=1)
            if title in figure_after:
                path, caption = figure_after[title]
                add_figure(doc, path, caption)
            if title == "9 示例结果与解释":
                add_result_table(doc)
            i += 1
            continue
        if line.startswith("### "):
            title = line[4:].strip()
            doc.add_heading(title, level=2)
            if title in figure_after:
                path, caption = figure_after[title]
                add_figure(doc, path, caption)
            i += 1
            continue
        if line.startswith("| "):
            raw_rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                raw_rows.append([x.strip() for x in lines[i].strip().strip("|").split("|")])
                i += 1
            data = [row for j, row in enumerate(raw_rows) if j != 1]
            table = doc.add_table(rows=1, cols=len(data[0]))
            table.style = "Table Grid"
            for c, value in enumerate(data[0]):
                cell = table.rows[0].cells[c]
                cell.text = value
                shade(cell._tc, PALE_BLUE)
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in cell.paragraphs[0].runs:
                    set_run_font(run, size=9, bold=True, color=BLUE)
            for row_values in data[1:]:
                cells = table.add_row().cells
                for c, value in enumerate(row_values):
                    cells[c].text = value
                    for run in cells[c].paragraphs[0].runs:
                        set_run_font(run, size=8.6)
            widths = [1700, 3260, 4400] if len(data[0]) == 3 else [TABLE_WIDTH // len(data[0])] * len(data[0])
            set_table_geometry(table, widths)
            doc.add_paragraph()
            continue
        if re.match(r"^\d+\. ", line):
            text = re.sub(r"^\d+\. ", "", line)
            if active_num_id is None:
                active_num_id = new_decimal_numbering(doc)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.375)
            p.paragraph_format.first_line_indent = Inches(-0.194)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.208
            apply_numbering(p, active_num_id)
            add_inline_markup(p, text)
            i += 1
            continue
        active_num_id = None
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline_markup(p, line[2:])
            i += 1
            continue
        if is_equation(line):
            p = doc.add_paragraph(style="Equation")
            run = p.add_run(line)
            set_run_font(run, latin="Cambria Math", east_asia="Noto Sans SC", size=10.5, color=BLUE)
            i += 1
            continue
        p = doc.add_paragraph()
        add_inline_markup(p, line)
        i += 1

    add_figure(doc, FIG / "ablation_cost.png", "图5  关键模块消融实验结果")
    add_figure(doc, ROOT / "outputs" / "northern" / "dispatch.png", "图6  北方示例典型日热源调度与基础热负荷")
    add_figure(doc, ROOT / "outputs" / "northern" / "network.png", "图7  北方示例最优管网建设方案")
    add_references(doc)

    settings = doc.settings._element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
