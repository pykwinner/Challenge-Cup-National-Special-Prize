from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from upgrade_algorithm_report_v11_1 import (
    add_paragraph_before,
    move_before,
    set_run_font,
    set_table_geometry,
    shade,
)
from upgrade_algorithm_report_v11_3 import align_table_indents


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "为什么数学建模能够支撑供热转型决策_算法增强版V5.docx"
OUTPUT = ROOT / "为什么数学建模能够支撑供热转型决策_算法增强版V6.docx"
DASHBOARD = ROOT / "outputs" / "v11_4_figures" / "representative_replay_closure_dashboard.png"


def refresh_summary(doc) -> None:
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            if "为什么可信：V11.3" in paragraph.text and "V11.3" in run.text:
                run.text = run.text.replace("V11.3", "V11.4")
            if "全工程58项" in paragraph.text and "58" in run.text:
                run.text = run.text.replace("58", "61")


def add_table_before(doc, anchor, headers, rows, widths) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, value in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = value
        shade(cell, "E8EEF5")
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in cell.paragraphs[0].runs:
            set_run_font(run, size=8.9, bold=True, color="17324D")
    for row_data in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_data):
            cells[index].text = str(value)
            cells[index].paragraphs[0].alignment = (
                WD_ALIGN_PARAGRAPH.LEFT if index == len(cells) - 1 else WD_ALIGN_PARAGRAPH.CENTER
            )
            for run in cells[index].paragraphs[0].runs:
                set_run_font(run, size=8.5)
    set_table_geometry(table, widths)
    anchor._p.addprevious(table._tbl)


def add_v11_4_section(doc) -> None:
    anchor = next(
        paragraph
        for paragraph in doc.paragraphs
        if paragraph.text.strip().startswith("3  为什么该模型具有工程物理可信性")
    )
    add_paragraph_before(
        doc,
        anchor,
        "2.7.8  V11.4：代表时段闭环修复与可靠容量预审",
        "Heading 3",
    )
    add_paragraph_before(
        doc,
        anchor,
        "V11.4使用乌鲁木齐、固原和武汉三套8760小时公共数据，形成“批量准入—逐序列极端日—设备最坏性能日—受约束权重校准—投资稳定性—固定投资全年回放”的闭环。代表日不再只锚定总负荷：每个非零热/冷负荷列分别保留最大峰值和最大爬坡日；空气源热泵再按低温、高湿除霜和停机边界逐时计算最坏1/COP日。相同日期按原日索引去重，新增极端日从所属聚类扣除1天，避免重复年化。",
        "Normal",
    )
    equation = add_paragraph_before(
        doc,
        anchor,
        "min Σⱼ[(Σd Aⱼd w_d − Eⱼʸ)/max(|Eⱼʸ|,1)]² + λΣd[(w_d−w⁰_d)/max(w⁰_d,1)]²",
        "Normal",
        WD_ALIGN_PARAGRAPH.CENTER,
    )
    equation.paragraph_format.keep_together = True
    add_paragraph_before(
        doc,
        anchor,
        "权重校准同时满足Σdw_d=365、w_d≥0、极端日w_d≥1；场景概率只进入Aⱼd核算而不被修改。算法只生成候选CSV和审计表，不覆盖原文件。连续年度只有一个实现场景时，规划期base等政策场景过滤器可显式映射到full_year并留痕；多场景名称不一致时仍然拒绝自动映射。",
        "Normal",
    )

    picture = doc.add_paragraph()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.paragraph_format.keep_with_next = True
    shape = picture.add_run().add_picture(str(DASHBOARD), width=Inches(6.25))
    shape._inline.docPr.set(
        "descr",
        "三场景代表日数量变化、武汉爬坡与年能量修复、配电缺口消除及8760小时回放结果",
    )
    shape._inline.docPr.set("title", "V11.4代表日修复与武汉全年回放闭环")
    move_before(picture, anchor)
    caption = add_paragraph_before(
        doc,
        anchor,
        "图 9  代表日风险修复与武汉8760小时固定投资回放闭环",
        "Caption" if "Caption" in doc.styles else "Normal",
        WD_ALIGN_PARAGRAPH.CENTER,
    )
    caption.paragraph_format.space_after = Pt(8)

    add_table_before(
        doc,
        anchor,
        ["指标", "原武汉方案", "V11.4候选", "工程解释"],
        [
            ("代表日", "9", "14", "新增逐负荷爬坡与最坏COP日"),
            ("最小爬坡捕获", "84.7%", "100%", "两个原告警消除"),
            ("最大热负荷年能量偏差", "6.31%", "0.020%", "365天权重守恒"),
            ("可靠配电容量", "551.613 kW", "573.855 kW", "覆盖全年最坏电热比"),
            ("全年配电缺口", "−0.917 kW", "约0 kW", "旧方案构模前终止"),
            ("8760小时回放", "未通过预审", "状态0、热冷零失供", "固定11类投资，不补建"),
        ],
        [1900, 1900, 2100, 3460],
    )

    add_paragraph_before(doc, anchor, "2.7.9  可靠容量审计与全年实测", "Heading 3")
    add_paragraph_before(
        doc,
        anchor,
        "大型全年MILP构建之前，程序按完整年度复算有效热峰值、12%容量裕度、建筑可靠柔性、热源最坏可用率、储热功率—能量—持续时间信用，以及portfolio_dispatch下的最低可靠电力需求。旧方案热侧缺口为0，但配电可靠容量少0.917kW，因此保存规划方案、冻结配置和逐设备审计后提前终止。该预审不替代逐时调度，只删除已知必然不可行的大模型构建。",
        "Normal",
    )
    equation = add_paragraph_before(
        doc,
        anchor,
        "ΔPᶠⁱʳᵐ = Pᵍʳⁱᵈ + Pˡᵒᶜᵃˡ,ᶠⁱʳᵐ − min Σs eˢᶠⁱʳᵐ qˢᶠⁱʳᵐ ≥ 0",
        "Normal",
        WD_ALIGN_PARAGRAPH.CENTER,
    )
    equation.paragraph_format.keep_together = True
    add_table_before(
        doc,
        anchor,
        ["场景", "代表日变化", "最大能量偏差", "最小爬坡", "候选审计"],
        [
            ("乌鲁木齐", "9→8", "0.423%→0.044%", "100%→100%", "0告警"),
            ("固原", "9→10", "0.494%→0.012%", "100%→100%", "0告警"),
            ("武汉", "9→14", "6.314%→0.020%", "84.731%→100%", "2→0告警"),
        ],
        [1700, 1650, 2200, 1900, 1910],
    )
    add_paragraph_before(
        doc,
        anchor,
        "武汉候选方案的全年回放含313579个变量、26280个整数变量和343761条约束；构模2.04s、求解44.11s。规划与回放状态均为0，热/冷失供率均为0，11类投资全部冻结，两座PCM首末SOC漂移均为0，地源热提取/排放比0.961位于0.80—1.25约束区间。全工程61项自动化测试通过。",
        "Normal",
    )


def main() -> None:
    for path in (SOURCE, DASHBOARD):
        if not path.exists():
            raise FileNotFoundError(path)
    doc = Document(SOURCE)
    refresh_summary(doc)
    add_v11_4_section(doc)
    align_table_indents(doc)
    doc.core_properties.title = "为什么数学建模能够支撑供热转型决策——算法增强版V6"
    doc.core_properties.subject = "V11.4代表时段闭环修复与固定投资可靠容量预审"
    doc.core_properties.version = "11.4"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
