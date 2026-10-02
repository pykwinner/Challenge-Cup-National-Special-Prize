from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from upgrade_algorithm_report_v11_1 import (
    add_paragraph_before,
    move_before,
)
from upgrade_algorithm_report_v11_3 import align_table_indents
from upgrade_algorithm_report_v11_4 import add_table_before


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "为什么数学建模能够支撑供热转型决策_算法增强版V6.docx"
OUTPUT = ROOT / "为什么数学建模能够支撑供热转型决策_算法增强版V7.docx"
DASHBOARD = (
    ROOT / "outputs" / "v11_5_figures" / "three_scenario_annual_replay_dashboard.png"
)
MATRIX = ROOT / "outputs" / "v11_5_figures" / "annual_replay_acceptance_matrix.png"


def refresh_summary(doc) -> None:
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            if "为什么可信：V11.4" in paragraph.text and "V11.4" in run.text:
                run.text = run.text.replace("V11.4", "V11.5")
            if "全工程61项" in paragraph.text and "61" in run.text:
                run.text = run.text.replace("61", "68")


def add_picture_before(doc, anchor, path, *, width, description, title, caption):
    picture = doc.add_paragraph()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.paragraph_format.keep_with_next = True
    shape = picture.add_run().add_picture(str(path), width=Inches(width))
    shape._inline.docPr.set("descr", description)
    shape._inline.docPr.set("title", title)
    move_before(picture, anchor)
    caption_paragraph = add_paragraph_before(
        doc,
        anchor,
        caption,
        "Caption" if "Caption" in doc.styles else "Normal",
        WD_ALIGN_PARAGRAPH.CENTER,
    )
    caption_paragraph.paragraph_format.space_after = Pt(8)


def add_v11_5_section(doc) -> None:
    anchor = next(
        paragraph
        for paragraph in doc.paragraphs
        if paragraph.text.strip().startswith("3  为什么该模型具有工程物理可信性")
    )
    add_paragraph_before(
        doc,
        anchor,
        "2.7.10  V11.5：三类场景固定投资全年回放",
        "Heading 3",
    )
    add_paragraph_before(
        doc,
        anchor,
        "V11.5把V11.4的武汉单场景闭环扩展到赛题三类典型变革场景。每个场景先用候选代表日决定热源、储热、管网、配电和末端改造投资，再冻结全部资产和候选余量，以连续8760小时公共原型数据重新优化运行。三场景按清单隔离启动，输入文件字节和关键参数形成SHA-256指纹；只有指纹一致且既有结果已认证通过时才断点续跑，防止配置、负荷或MIP gap变化后误用旧证据。",
        "Normal",
    )
    add_picture_before(
        doc,
        anchor,
        DASHBOARD,
        width=6.35,
        description="乌鲁木齐、固原和武汉三类场景的全年模型规模、求解时间、零失供与验收门结果",
        title="V11.5三场景全年固定投资回放证据",
        caption="图 10  三类供热场景8760小时固定投资回放结果",
    )
    add_table_before(
        doc,
        anchor,
        ["赛题场景", "城市原型", "变量", "整数变量", "求解/s", "验收"],
        [
            ("大型集中供热低碳重构", "乌鲁木齐", "449,335", "26,283", "83.74", "10/10"),
            ("分散式独立供热清洁转型", "固原", "423,898", "26,280", "45.89", "10/10"),
            ("夏热冬冷高品质系统新建", "武汉", "313,579", "26,280", "55.33", "12/12"),
        ],
        [2600, 1300, 1450, 1450, 1200, 1360],
    )
    add_paragraph_before(
        doc,
        anchor,
        "三场景代表期规划和全年回放求解状态均为0，全部投资冻结，年度与最差场景热失供率均为0；武汉冷失供率亦为0。三场景合计1186812个变量、78843个整数变量和1334272条约束。S1可靠热容量裕度1100kW；S2和S3热侧与配电可靠容量贴边满足，提示极端负荷或设备性能改变后必须重新回放。",
        "Normal",
    )

    add_paragraph_before(
        doc,
        anchor,
        "2.7.11  机器可读验收证书与储热边界口径修正",
        "Heading 3",
    )
    add_paragraph_before(
        doc,
        anchor,
        "每次正式回放生成annual_replay_acceptance_certificate.json。通用验收门逐项检查完整公历年、代表时段覆盖、规划与回放最优、投资冻结、可靠容量、年度及最差场景失供、储热边界和清洁热约束；武汉另检查供冷可靠性与地源年度热平衡。只有全部适用门通过，场景才获得认证，而不是把单一solver status当成全部结论。",
        "Normal",
    )
    equation = add_paragraph_before(
        doc,
        anchor,
        "Zᵃⁿⁿᵘᵃˡ = Πₖ zₖ；仅当全部适用 zₖ=1 时，certification_pass=true",
        "Normal",
        WD_ALIGN_PARAGRAPH.CENTER,
    )
    equation.paragraph_format.keep_together = True
    add_picture_before(
        doc,
        anchor,
        MATRIX,
        width=6.35,
        description="三个供热场景在连续年度、覆盖、求解、投资冻结、容量、失供、储热、供冷、地源热平衡和清洁热等验收门的通过矩阵",
        title="V11.5全年回放验收矩阵",
        caption="图 11  三类场景全年回放机器可读验收矩阵",
    )
    add_paragraph_before(
        doc,
        anchor,
        "储热SOC输出表示每个时段结束后的状态。旧后检查直接用末时段SOC减首时段SOC，会把首小时充放热误报为跨年漂移。V11.5在调度表中显式输出boundary_group；循环边界把末时段SOC作为首时段之前的状态，并按真实状态方程重建数值残差：",
        "Normal",
    )
    equation = add_paragraph_before(
        doc,
        anchor,
        "rᶜʸᶜ = Eₜ₀ − λEₜᴺ − ηᶜʰQᶜʰₜ₀Δt + Qᵈⁱˢₜ₀Δt/ηᵈⁱˢ",
        "Normal",
        WD_ALIGN_PARAGRAPH.CENTER,
    )
    equation.paragraph_format.keep_together = True
    add_table_before(
        doc,
        anchor,
        ["后检查对象", "旧显示", "V11.5结果", "判定"],
        [
            ("S1低温水罐循环边界", "18.909 kWh伪漂移", "漂移0；残差−7.1×10⁻¹⁵ kWh", "通过"),
            ("S3两座PCM循环边界", "首末均为0", "漂移0；残差0", "通过"),
            ("武汉地源年度平衡", "取热/排热0.961", "取热/排热0.957", "0.80—1.25内"),
        ],
        [2450, 2200, 3000, 1710],
    )
    add_paragraph_before(
        doc,
        anchor,
        "全工程68项自动化测试通过。代表日模型五次重复构模中位数为S1 0.0632s、S2 0.0508s、S3 0.0446s；MILP结构与V11.4一致，新增成本位于批量编排、结果后检查和证书生成层。公开数据仍是天气与规则化负荷代理，正式项目必须替换为热表、建筑、设备、管网和配电实测后重新生成代表日并回放。",
        "Normal",
    )


def main() -> None:
    for path in (SOURCE, DASHBOARD, MATRIX):
        if not path.exists():
            raise FileNotFoundError(path)
    doc = Document(SOURCE)
    refresh_summary(doc)
    add_v11_5_section(doc)
    align_table_indents(doc)
    doc.core_properties.title = "为什么数学建模能够支撑供热转型决策——算法增强版V7"
    doc.core_properties.subject = "V11.5三场景固定投资全年回放与机器可读验收"
    doc.core_properties.version = "11.5"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
