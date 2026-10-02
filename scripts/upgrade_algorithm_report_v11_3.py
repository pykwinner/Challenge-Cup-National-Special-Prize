from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from upgrade_algorithm_report_v11_1 import (
    add_paragraph_before,
    move_before,
    set_run_font,
    set_table_geometry,
    shade,
)


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "为什么数学建模能够支撑供热转型决策_算法增强版V4.docx"
OUTPUT = ROOT / "为什么数学建模能够支撑供热转型决策_算法增强版V5.docx"
DASHBOARD = ROOT / "outputs" / "v11_3_figures" / "model_size_runtime_dashboard.png"


def align_table_indents(doc) -> None:
    """Align visible table borders with their first-cell text margins."""

    for table in doc.tables:
        tc_pr = table.rows[0].cells[0]._tc.get_or_add_tcPr()
        margins = tc_pr.find(qn("w:tcMar"))
        start = None if margins is None else margins.find(qn("w:start"))
        if start is None:
            continue
        width = start.get(qn("w:w"))
        if width is None:
            continue
        tbl_pr = table._tbl.tblPr
        indent = tbl_pr.find(qn("w:tblInd"))
        if indent is None:
            indent = OxmlElement("w:tblInd")
            tbl_pr.append(indent)
        indent.set(qn("w:type"), "dxa")
        indent.set(qn("w:w"), width)


def refresh_version_summary(doc) -> None:
    for paragraph in doc.paragraphs:
        if "为什么可信：V11.2" in paragraph.text:
            for run in paragraph.runs:
                if "V11.2" in run.text:
                    run.text = run.text.replace("V11.2", "V11.3")
            break
    for paragraph in doc.paragraphs:
        if "全工程57项自动化测试通过" in paragraph.text:
            for run in paragraph.runs:
                run.text = run.text.replace("全工程57项", "全工程58项")
                run.text = run.text.replace(
                    "三个公共代表日场景均保持最优与零失供",
                    "三个公共场景在显式1% gap口径下均以状态0结束且零失供",
                )
    benchmark_updates = {
        "C1 乌鲁木齐": ("0.0657 s", "1% gap口径0.875%；零失供"),
        "C2 固原": ("0.0502 s", "1% gap口径0.088%；零失供"),
        "C3 武汉": ("0.0458 s", "1% gap口径0.600%；零失供"),
    }
    for table in doc.tables:
        for row in table.rows[1:]:
            key = row.cells[0].text.strip()
            if key not in benchmark_updates or len(row.cells) < 4:
                continue
            build_time, evidence = benchmark_updates[key]
            row.cells[2].text = build_time
            row.cells[3].text = evidence
            for cell in (row.cells[2], row.cells[3]):
                for run in cell.paragraphs[0].runs:
                    set_run_font(run, size=9.0)


def add_v11_3_section(doc) -> None:
    anchor = next(
        paragraph
        for paragraph in doc.paragraphs
        if paragraph.text.strip().startswith("3  为什么该模型具有工程物理可信性")
    )
    add_paragraph_before(
        doc,
        anchor,
        "2.7.6  V11.3：全年规模先验、运行画像与安全修正建议",
        "Heading 3",
    )
    add_paragraph_before(
        doc,
        anchor,
        "V11.3在正式8760/8784小时求解之前，用两个短前缀只构模、不求解，识别变量数、整数变量数、约束数、稀疏矩阵非零元和内存代理。变量、约束与非零元采用双样本仿射投影；构模时间只给宽区间。MILP求解时间明确返回空值，因为它受预处理、整数松弛质量、初始可行解、分支定界、割平面、MIP gap和时间上限共同影响，不能由矩阵规模线性外推。",
        "Normal",
    )
    equation = add_paragraph_before(
        doc,
        anchor,
        "Nₖ(T) ≈ αₖ + βₖT，k∈{变量、整数变量、约束、非零元}；tˢᵒˡᵛᵉ(T) 不作线性外推",
        "Normal",
        WD_ALIGN_PARAGRAPH.CENTER,
    )
    equation.paragraph_format.keep_together = True

    picture = doc.add_paragraph()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.paragraph_format.keep_with_next = True
    shape = picture.add_run().add_picture(str(DASHBOARD), width=Inches(6.25))
    shape._inline.docPr.set(
        "descr",
        "8760小时变量约束非零元和内存代理，以及168小时规划回放构模求解耗时看板",
    )
    shape._inline.docPr.set("title", "V11.3全年模型规模与运行画像")
    move_before(picture, anchor)
    caption = add_paragraph_before(
        doc,
        anchor,
        "图 8  全年模型规模先验与168小时固定投资回放实测",
        "Caption" if "Caption" in doc.styles else "Normal",
        WD_ALIGN_PARAGRAPH.CENTER,
    )
    caption.paragraph_format.space_after = Pt(8)

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = ["审计量", "8760小时先验", "168小时校验", "报告边界"]
    for index, value in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = value
        shade(cell, "E8EEF5")
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in cell.paragraphs[0].runs:
            set_run_font(run, size=9.0, bold=True, color="17324D")
    rows = [
        ("变量/整数变量", "446785 / 52566", "投影8593，实测8581", "整数规模仅用于风险分级"),
        ("约束/非零元", "461382 / 1743303", "约束误差0.02%", "按当前配置与数据列生成"),
        ("内存代理", "CSR+向量40.7 MB；构模313.3 MB", "回放构模代理5.83 MB", "不含求解器树、割与系统开销"),
        ("构模时间", "中值2.14 s；宽区间1.06–6.45 s", "投影0.042 s，实测0.044 s", "只估构模，不等于求解"),
        ("MILP求解时间", "不估计", "规划0.309 s；回放0.575 s", "正式年度须限时实测并报告gap"),
    ]
    for row_data in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_data):
            cells[index].text = value
            cells[index].paragraphs[0].alignment = (
                WD_ALIGN_PARAGRAPH.CENTER if index != 3 else WD_ALIGN_PARAGRAPH.LEFT
            )
            for run in cells[index].paragraphs[0].runs:
                set_run_font(run, size=8.6)
    set_table_geometry(table, [1720, 2550, 2360, 2730])
    anchor._p.addprevious(table._tbl)

    add_paragraph_before(doc, anchor, "2.7.7  代表时段修正的安全边界", "Heading 3")
    add_paragraph_before(
        doc,
        anchor,
        "覆盖审计把问题翻译成机器可读建议：峰值缺失则强制最大负荷日，爬坡缺失则强制最大爬坡日，低温或资源低谷缺失则补充对应极端日，年能量偏差超过10%则建议重新校准权重或增加季节代表日。系统同时给出Eʸ/Eᴿ作为诊断性倍率，但automatic_application_allowed固定为false。原因是直接乘权重可能破坏代表日总天数、场景概率和既有极端锚点；正式修正必须重新聚类或在保持这些约束下校准。",
        "Normal",
    )
    add_paragraph_before(
        doc,
        anchor,
        "当前合成年度的四个负荷列均建议倍率约1.166，表示原代表时段折算能量偏低约14.2%；该倍率用于定位问题，不会被程序静默写回配置。168小时固定投资回放中，规划与回放均达到最优且零失供，全部投资冻结审计通过。",
        "Normal",
    )


def main() -> None:
    for path in (SOURCE, DASHBOARD):
        if not path.exists():
            raise FileNotFoundError(path)
    doc = Document(SOURCE)
    refresh_version_summary(doc)
    add_v11_3_section(doc)
    align_table_indents(doc)
    doc.core_properties.title = "为什么数学建模能够支撑供热转型决策——算法增强版V5"
    doc.core_properties.subject = "V11.3全年模型规模先验、运行画像与安全修正建议"
    doc.core_properties.version = "11.3"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
