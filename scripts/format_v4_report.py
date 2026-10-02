from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


BLUE = "17324D"
MID_BLUE = "2E74B5"
PALE_BLUE = "F4F6F9"
LIGHT_GREY = "F4F6F8"
FONT = "Noto Sans SC"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, size: float | None = None, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = FONT
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{key}"), FONT)


def set_table_widths(table, widths: list[int]) -> None:
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), "120")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    cell_mar = tbl_pr.find(qn("w:tblCellMar"))
    if cell_mar is None:
        cell_mar = OxmlElement("w:tblCellMar")
        tbl_pr.append(cell_mar)
    for side, value in (("top", 80), ("bottom", 80), ("start", 120), ("end", 120)):
        element = cell_mar.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            cell_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        row._tr.get_or_add_trPr()
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, 8.5, color="667788")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)
    run = paragraph.add_run(" 页")
    set_run_font(run, 8.5, color="667788")


def format_report(path: Path) -> None:
    doc = Document(path)
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.9)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(11)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.line_spacing = 1.333
    normal.paragraph_format.space_after = Pt(8)

    for style_name in ("Body Text", "List Paragraph", "Block Text"):
        if style_name not in doc.styles:
            continue
        style = doc.styles[style_name]
        style.font.name = FONT
        style.font.size = Pt(11)
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
        style.paragraph_format.line_spacing = 1.333
        style.paragraph_format.space_after = Pt(8)

    if "List Paragraph" in doc.styles:
        list_style = doc.styles["List Paragraph"]
        list_style.paragraph_format.left_indent = Cm(0.9525)
        list_style.paragraph_format.first_line_indent = Cm(-0.4928)
        list_style.paragraph_format.line_spacing = 1.208
        list_style.paragraph_format.space_after = Pt(4)

    heading_specs = {
        "Title": (23, BLUE, True, 4, 8),
        "Subtitle": (12.5, MID_BLUE, False, 0, 14),
        "Heading 1": (16, MID_BLUE, True, 18, 10),
        "Heading 2": (13, MID_BLUE, True, 12, 6),
        "Heading 3": (12, BLUE, True, 8, 4),
        "Heading 4": (10.8, BLUE, True, 7, 4),
    }
    for name, (size, color, bold, before, after) in heading_specs.items():
        if name not in doc.styles:
            continue
        style = doc.styles[name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    # Turn the first two Markdown headings into a compact report title block.
    if doc.paragraphs:
        doc.paragraphs[0].style = doc.styles["Title"]
        doc.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if len(doc.paragraphs) > 1:
        doc.paragraphs[1].style = doc.styles["Subtitle"]
        doc.paragraphs[1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if len(doc.paragraphs) > 2:
        doc.paragraphs[2].alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.paragraphs[2].paragraph_format.line_spacing = 1.45
        doc.paragraphs[2].paragraph_format.space_before = Pt(6)
        doc.paragraphs[2].paragraph_format.space_after = Pt(12)

    for paragraph in doc.paragraphs:
        if paragraph.style.name in {"Normal", "Body Text", "Block Text"}:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if paragraph.text.startswith("重要说明："):
            paragraph.paragraph_format.left_indent = Cm(0.35)
            paragraph.paragraph_format.right_indent = Cm(0.35)
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(10)
            p_pr = paragraph._p.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), PALE_BLUE)
            p_pr.append(shd)
        for run in paragraph.runs:
            set_run_font(run)

    width_sets = {
        2: [2600, 6760],
        3: [2200, 3900, 3260],
        4: [1900, 2350, 3000, 2110],
        5: [1750, 1300, 1850, 1900, 2560],
        6: [1650, 900, 1750, 1250, 1800, 2010],
    }
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid" if "Table Grid" in doc.styles else "Table"
        widths = width_sets.get(len(table.columns), [9360 // len(table.columns)] * len(table.columns))
        set_table_widths(table, widths)
        set_repeat_table_header(table.rows[0])
        for row_idx, row in enumerate(table.rows):
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if row_idx == 0:
                    set_cell_shading(cell, PALE_BLUE)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(2)
                    paragraph.paragraph_format.line_spacing = 1.05
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    for run in paragraph.runs:
                        set_run_font(run, size=8.6, bold=(row_idx == 0), color=(BLUE if row_idx == 0 else None))

    for paragraph in doc.paragraphs:
        if paragraph.style.name in {"Display Math", "Equation"}:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(6)
            paragraph.paragraph_format.keep_together = True

    version = next((candidate for candidate in ("V10", "V9", "V8", "V7", "V6", "V5", "V4") if candidate in path.stem), "V4")
    header = section.header.paragraphs[0]
    header.text = f"NHR-TIP｜北方城镇供热 {version} 韧性转型模型"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        set_run_font(run, size=8.5, color="667788")
    footer = section.footer.paragraphs[0]
    footer.clear()
    add_page_number(footer)

    doc.core_properties.title = path.stem
    doc.core_properties.subject = "北方供热源网荷储、电网容量、温级改造、复合寒潮与N-1韧性规划"
    doc.core_properties.keywords = "北方供热; 混合整数线性规划; N-1; 热泵; 储热; 低温热网"
    doc.core_properties.description = "narrative_proposal preset with named A4 and Chinese-font overrides"
    doc.save(path)


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("usage: format_v4_report.py REPORT.docx")
    format_report(Path(sys.argv[1]).resolve())
