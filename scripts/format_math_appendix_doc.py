from __future__ import annotations

from pathlib import Path
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from build_northern_model_doc import set_run_font, set_table_geometry, shade


ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "docs" / "北方供热优化数学模型与公式附录.docx"
BLUE = "17324D"
MID_BLUE = "2E74B5"
PALE_BLUE = "E8EEF5"


def format_document(path: Path = DOCX) -> None:
    doc = Document(path)

    def set_east_asia(run, family: str = "Noto Sans SC") -> None:
        r_pr = run._element.get_or_add_rPr()
        r_fonts = r_pr.rFonts
        if r_fonts is None:
            r_fonts = OxmlElement("w:rFonts")
            r_pr.insert(0, r_fonts)
        r_fonts.set(qn("w:eastAsia"), family)

    style_specs = {
        "Title": (24, BLUE, True, 4, 12),
        "Subtitle": (12, MID_BLUE, False, 0, 18),
        "Heading 1": (16, MID_BLUE, True, 16, 8),
        "Heading 2": (13, MID_BLUE, True, 10, 5),
        "Heading 3": (11.5, BLUE, True, 8, 4),
    }
    for name, (size, color, bold, before, after) in style_specs.items():
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Noto Sans SC")
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = name.startswith("Heading")

    for paragraph in doc.paragraphs:
        if paragraph.style.name in style_specs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in paragraph.runs:
            set_east_asia(run)

    width_sets = {
        2: [2200, 7160],
        3: [2300, 4200, 2860],
    }
    for table in doc.tables:
        table.style = "Table Grid"
        set_table_geometry(table, width_sets[len(table.columns)])
        for cell in table.rows[0].cells:
            shade(cell._tc, PALE_BLUE)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    set_run_font(run, size=9.2, color=BLUE, bold=True)
                    set_east_asia(run)
        for row in table.rows[1:]:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(2)
                    paragraph.paragraph_format.line_spacing = 1.05
                    for run in paragraph.runs:
                        set_run_font(run, size=9.0)
                        set_east_asia(run)

    # Make equations easier to scan and prevent formula rows from breaking awkwardly.
    for paragraph in doc.paragraphs:
        if paragraph.style.name in {"Display Math", "Equation"}:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(7)
            paragraph.paragraph_format.keep_together = True

    doc.core_properties.title = "北方城镇供热优化数学模型与公式附录"
    doc.core_properties.subject = "两阶段场景鲁棒MILP、可靠性、CVaR与线性化公式"
    doc.save(path)


if __name__ == "__main__":
    format_document(Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else DOCX)
