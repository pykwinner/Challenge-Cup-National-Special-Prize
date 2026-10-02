from __future__ import annotations

import json
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "competition_report_20260827"
FIG = OUT / "figures"
DOCX = OUT / "暖城向新_三类供热转型通用优化模型_竞赛研究报告.docx"
CSV = OUT / "仿真结果数据总表.csv"

BLUE = "17365D"
CYAN = "2A7F9E"
ORANGE = "D97A2B"
LIGHT = "EAF2F8"
GRAY = "5B6573"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, *, bold=False, color="000000", size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        set_cell_shading(table.rows[0].cells[i], BLUE)
        set_cell_text(table.rows[0].cells[i], h, bold=True, color="FFFFFF", size=8.5)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            if ridx % 2 == 1:
                set_cell_shading(cells[i], "F3F6F8")
            set_cell_text(cells[i], value, size=8.2)
    if widths:
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def add_para(doc, text, *, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing = 1.35
    p.paragraph_format.space_after = Pt(4)
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def add_equation(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    r.font.name = "Cambria Math"
    r.font.size = Pt(10.5)
    return p


def add_figure(doc, path, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Cm(16.2))
    c = doc.add_paragraph(caption)
    c.style = "Caption" if "Caption" in [s.name for s in doc.styles] else "Normal"
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(7)


def load_data():
    batch = pd.read_csv(ROOT / "outputs/v11_5_three_scenario_annual_replay/annual_replay_batch_summary.csv")
    risk = pd.read_csv(ROOT / "outputs/v11_10_cvar_weight_scan/risk_weight_scan_summary.csv")
    multi = pd.read_csv(ROOT / "outputs/v11_9_multiyear_formal_replay/multiyear_replay_summary.csv")
    perf = {
        "baseline_2021": 575.746,
        "tight_2021": 86.903,
        "baseline_2024": 26.736,
        "tight_2024": 25.847,
    }
    return batch, risk, multi, perf


def setup_plot():
    plt.rcParams.update({
        "font.family": "sans-serif",
        "font.sans-serif": ["Microsoft YaHei", "SimHei", "Arial Unicode MS"],
        "axes.unicode_minus": False,
        "font.size": 11,
        "axes.titlesize": 14,
        "axes.labelsize": 11.5,
        "xtick.labelsize": 10,
        "ytick.labelsize": 10,
        "legend.fontsize": 9.5,
        "figure.dpi": 180,
        "savefig.dpi": 320,
    })


def savefig(fig, name):
    path = FIG / name
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def make_figures(batch, risk, multi, perf):
    setup_plot()
    FIG.mkdir(parents=True, exist_ok=True)
    figures = {}

    fig, ax = plt.subplots(figsize=(10, 5.2))
    ax.axis("off")
    boxes = [
        (0.02, .60, .18, .25, "数据与边界", "气象·负荷\n资源·政策"),
        (.24, .60, .20, .25, "城市层", "投资组合\n总量与风险"),
        (.48, .60, .20, .25, "区域层", "管网·微网\n储热·互联"),
        (.72, .60, .20, .25, "用户层", "室温柔性\n需求响应"),
    ]
    for x, y, w, h, title, sub in boxes:
        ax.add_patch(plt.Rectangle((x, y), w, h, facecolor="#EAF2F8", edgecolor="#17365D", lw=1.7))
        ax.text(x+w/2, y+h*.68, title, ha="center", va="center", weight="bold", color="#17365D", fontsize=13)
        ax.text(x+w/2, y+h*.31, sub, ha="center", va="center", color="#465466", fontsize=10.5)
    for x1, x2 in [(0.20, .24), (.44, .48), (.68, .72)]:
        ax.annotate("", xy=(x2, .725), xytext=(x1, .725), arrowprops=dict(arrowstyle="->", lw=1.8, color="#D97A2B"))
    ax.add_patch(plt.Rectangle((.16, .15), .68, .25, facecolor="#17365D", edgecolor="#17365D"))
    ax.text(.50, .33, "统一决策闭环", ha="center", va="center", color="white", weight="bold", fontsize=15)
    ax.text(.50, .235, "代表时段规划  →  投资冻结  →  连续逐时回放  →  极端年/CVaR认证", ha="center", va="center", color="white", fontsize=11.5)
    ax.text(.50, .08, "反馈：失供、容量裕度、碳排放与风险成本返回规划层", ha="center", va="center", color="#17365D", fontsize=10.5)
    figures["architecture"] = savefig(fig, "fig01_multiscale_closed_loop.png")

    profiles = ["scenario1_urumqi", "scenario2_guyuan", "scenario3_wuhan"]
    cities = ["乌鲁木齐", "固原", "武汉"]
    fig, axes = plt.subplots(1, 3, figsize=(10.5, 4.3), sharey=False)
    for ax, p, city in zip(axes, profiles, cities):
        f = ROOT / f"data/public_scenario_inputs/{p}_full_year_2023.csv"
        d = pd.read_csv(f)
        load_cols = [c for c in d.columns if c.startswith("demand_") and "cool" not in c]
        load = d[load_cols].sum(axis=1).to_numpy()
        load = np.sort(load)[::-1]
        ax.plot(np.linspace(0, 100, len(load)), load, color="#17365D", lw=2.1)
        ax.fill_between(np.linspace(0, 100, len(load)), load, color="#8FC3D4", alpha=.35)
        ax.set_title(city, weight="bold")
        ax.set_xlabel("年内累计小时占比 / %")
        ax.grid(alpha=.22)
    axes[0].set_ylabel("热负荷 / kW")
    fig.suptitle("三场景负荷持续曲线（2023公共原型数据）", weight="bold")
    fig.tight_layout()
    figures["ldc"] = savefig(fig, "fig02_load_duration_curves.png")

    p = "scenario1_urumqi"
    dispatch = pd.read_csv(ROOT / f"outputs/v11_5_three_scenario_annual_replay/{p}/fixed_investment_annual_replay/dispatch.csv")
    src = dispatch[dispatch.category.eq("source_heat")].pivot_table(index="time_index", columns="asset", values="value", aggfunc="sum", fill_value=0)
    preferred = src.sum().sort_values(ascending=False).head(6).index
    daily = src[preferred].groupby(src.index // 24).sum() / 24
    fig, ax = plt.subplots(figsize=(10.5, 5.0))
    im = ax.imshow(daily.T, aspect="auto", cmap="YlGnBu", interpolation="nearest")
    ax.set_yticks(range(len(preferred)))
    ax.set_yticklabels([x.replace("_", " ") for x in preferred])
    ax.set_xlabel("日序（1–365）")
    ax.set_title("乌鲁木齐源—荷匹配热力图：日均供热功率", weight="bold")
    cb = fig.colorbar(im, ax=ax, pad=.015)
    cb.set_label("kW")
    fig.tight_layout()
    figures["heatmap"] = savefig(fig, "fig03_source_load_heatmap.png")

    fig, ax = plt.subplots(figsize=(9.6, 5.0))
    for p, city, marker in zip(profiles, cities, ["o", "s", "^"]):
        q = risk[risk.profile.eq(p)].sort_values("cvar_weight")
        x = q["expected_cost_before_cvar_yuan_per_year"] - q.iloc[0]["expected_cost_before_cvar_yuan_per_year"]
        y = q.iloc[0]["cvar_operating_risk_cost_yuan"] - q["cvar_operating_risk_cost_yuan"]
        ax.plot(x, y, marker=marker, ms=8, lw=2, label=city)
        for _, row in q.iterrows():
            ax.annotate(f"λ={row.cvar_weight:g}", (row.expected_cost_before_cvar_yuan_per_year-q.iloc[0].expected_cost_before_cvar_yuan_per_year, q.iloc[0].cvar_operating_risk_cost_yuan-row.cvar_operating_risk_cost_yuan), xytext=(4, 5), textcoords="offset points", fontsize=8.5)
    ax.set_xlabel("相对风险中性方案的期望年成本增量 / 元")
    ax.set_ylabel("尾部风险成本降低量 / 元")
    ax.set_title("成本—风险前沿与统一推荐点", weight="bold")
    ax.legend(frameon=False)
    ax.grid(alpha=.25)
    fig.tight_layout()
    figures["risk"] = savefig(fig, "fig04_cost_risk_frontier.png")

    fig, ax = plt.subplots(figsize=(9.6, 4.8))
    base = [perf["baseline_2021"], perf["baseline_2024"]]
    tight = [perf["tight_2021"], perf["tight_2024"]]
    x = np.arange(2); w=.34
    ax.bar(x-w/2, base, w, label="V11.10", color="#AAB7C4")
    ax.bar(x+w/2, tight, w, label="V11.11紧化模型", color="#2A7F9E")
    for i, (b, t) in enumerate(zip(base, tight)):
        ax.text(i+w/2, t+max(base)*.025, f"{b/t:.2f}×", ha="center", weight="bold", color="#17365D")
    ax.set_xticks(x, ["固原2021", "固原2024"])
    ax.set_ylabel("求解时间 / s")
    ax.set_title("等投资方案下的MILP性能对比", weight="bold")
    ax.legend(frameon=False)
    ax.grid(axis="y", alpha=.22)
    fig.tight_layout()
    figures["performance"] = savefig(fig, "fig05_solver_performance.png")

    return figures


def configure_doc(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.35)
    section.right_margin = Cm(2.15)
    normal = doc.styles["Normal"]
    normal.font.name = "SimSun"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    normal.font.size = Pt(10.5)
    for style_name, size, color in [("Title", 24, BLUE), ("Heading 1", 16, BLUE), ("Heading 2", 13, CYAN), ("Heading 3", 11, GRAY)]:
        st = doc.styles[style_name]
        st.font.name = "Microsoft YaHei"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(color)
        st.font.bold = True
    for sec in doc.sections:
        footer = sec.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.add_run("暖城向新·三类供热转型通用模型  |  ")
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        footer._p.append(fld)


def build_doc(batch, risk, multi, figures):
    doc = Document()
    configure_doc(doc)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(90)
    r = p.add_run("暖城向新·热力不息")
    r.bold = True; r.font.size = Pt(26); r.font.color.rgb = RGBColor.from_string(BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("面向三类供热转型场景的多尺度源荷协同规划与逐时调度")
    r.font.size = Pt(17); r.font.color.rgb = RGBColor.from_string(CYAN)
    doc.add_paragraph()
    add_table(doc, ["报告定位", "核心交付"], [
        ("挑战杯详细研究报告", "数学模型、求解逻辑、逐时仿真、经济与碳评价"),
        ("技术版本", "V11.11 · 2026年8月27日"),
        ("证据口径", "2023年8760 h三场景回放 + 2020–2024多气象年 + CVaR扫描"),
    ], [4.5, 11.5])
    doc.add_paragraph()
    p = doc.add_paragraph("参赛团队：____________________")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

    add_heading(doc, "摘要", 1)
    add_para(doc, "针对既有大型集中供热低碳重构、分散式独立供热清洁转型和夏热冬冷地区高品质系统新建三类任务，本研究建立一套共享数学内核、通过场景参数适配的源—网—储—荷协同模型。模型以城市投资、区域能流与用户热惰性为三个相互约束的尺度，统一处理热源容量、管网互联、储热状态、建筑室温和可靠性备用。求解采用“代表时段规划—投资冻结—连续逐时回放—多气象年风险认证”工程闭环，并用期望净成本与尾部风险成本的协同目标筛选方案。")
    add_para(doc, "三类公共原型均完成2023年连续8760 h固定投资回放，求解状态均为0，热失供率均为0，武汉冷失供率为0，32/32项适用验收门通过。七个极端气象年累计61392 h回放全部通过。在α=0.80的CVaR口径下，λ=0.05为三场景统一工程折中：乌鲁木齐和固原以很小期望成本增量降低尾部成本，武汉投资结构在扫描区间内保持不变。通过精确启停状态凸包和稀疏矩阵缓存，固原典型困难年的整数变量减少66.7%，等投资方案回放从575.746 s降至86.903 s。")
    add_para(doc, "关键词：供热转型；多尺度优化；源荷匹配；混合整数线性规划；连续逐时回放；CVaR；碳排放")

    add_heading(doc, "1 问题界定与评分对齐", 1)
    add_para(doc, "赛题的难点不在于为三座城市分别配置技术，而在于用同一决策框架回答三个不同的工程问题：既有资产如何留用与替代，薄弱基础设施如何从散烧跨越到集约供能，零基建地区如何建立按需响应的冷热一体系统。因此，本文不把技术路线预先固定，而把资源可用性、基建底数、气象、价格、排放因子和舒适度边界写入参数层，由统一模型选择容量与逐时调度。")
    add_table(doc, ["评分点", "正文证据", "机器可读交付"], [
        ("多目标协同与收敛", "式(9)–(11)、成本—风险前沿、12/12状态0解", "risk_weight_scan_summary.csv"),
        ("多尺度耦合", "城市投资—区域能流—用户热惰性联立", "dispatch.csv / investments.csv"),
        ("动态特性", "8760 h连续回放、储热SOC、室温状态、启停约束", "acceptance_certificate.json"),
        ("经济与碳评价", "年化投资+运行+碳+可靠性台账", "summary.json / 结果总表"),
        ("代码可移植性", "配置、模型、求解、回放、证书分层", "src/heatplan + tests"),
    ], [3.5, 7.4, 5.1])

    add_heading(doc, "2 数据与负荷预测模型", 1)
    add_heading(doc, "2.1 数据口径与供暖季范围", 2)
    add_para(doc, "公共原型数据以逐时室外温度、电价、电网排放因子、余热可用率、可再生电力和分区负荷为入口。为避免把全年能耗错当成赛题要求的供暖季，正式报告将两种证据分开：已完成的2023年8760 h全年回放用于证明模型的连续时序能力；2024–2025季的乌鲁木齐4392 h、固厘3984 h、武汉2880 h作为赛题专项回放范围，未执行前不写为“已通过”。")
    add_table(doc, ["城市", "专项时窗", "小时数", "当前状态"], [
        ("乌鲁木齐", "2024-10-10 00:00 至 2025-04-10 23:00", "4392", "待专项回放"),
        ("固原", "2024-11-01 00:00 至 2025-04-15 23:00", "3984", "待专项回放"),
        ("武汉", "2024-11-16 00:00 至 2025-03-15 23:00", "2880", "待专项回放"),
    ], [3, 8.2, 2, 2.8])
    add_heading(doc, "2.2 负荷推导", 2)
    add_para(doc, "负荷预测采用物理可解释的基荷与一阶热惯性表达。对分区n，以建筑传热系数UAₙ、设计室温Tⁱᵐᵗ、室外温度Tᵒᵘᵗ和占用/日历因子φₙ,ₜ构造无柔性基荷：")
    add_equation(doc, "Q⁰ₙ,ₜ = max{0, UAₙ(Tⁱᵐᵗₙ − Tᵒᵘᵗ,ₜ) + Qᵛⁿᵗₙ,ₜ}·φₙ,ₜ")
    add_para(doc, "为表示建筑蓄热与需求响应，引入等效室温状态θₙ,ₜ。在时步Δt下：")
    add_equation(doc, "Cₙ(θₙ,ₜ − θₙ,ₜ₋₁) = Δt[Qₙ,ₜ − UAₙ(θₙ,ₜ − Tᵒᵘᵗ,ₜ) − Qᵛⁿᵗₙ,ₜ]")
    add_para(doc, "其中θ̅ₙ≤θₙ,ₜ≤θ̄ₙ把舒适度转化为显式边界。数据预处理对缺失、闰年、单位和时区进行校验；代表时段仅用于投资筛选，最终运行可行性必须在原始时序上复核。")
    add_figure(doc, figures["ldc"], "图1  三类场景负荷持续曲线（大字号、统一坐标语义）")

    add_heading(doc, "3 多尺度源荷协同模型", 1)
    add_figure(doc, figures["architecture"], "图2  城市—区域—用户三尺度决策闭环")
    add_heading(doc, "3.1 变量与能量平衡", 2)
    add_para(doc, "城市层决定热源、储热、管网和配电容量x；区域层决定热源出力qᵍ,ₜ、管道流量fₑ,ₜ和储热充放qᶜʰ/qᵈʲˢ；用户层以室温状态与可调负荷表示舒适度。对任意节点n、品位g和时刻t，热平衡为：")
    add_equation(doc, "Σᵢ qᵢ,ₙ,ᵍ,ₜ + Σₑ∈δ⁻(n) ηₑfₑ,ᵍ,ₜ + qᵈʲˢₙ,ₜ + uₙ,ₜ = Qₙ,ᵍ,ₜ + Σₑ∈δ⁺(n)fₑ,ᵍ,ₜ + qᶜʰₙ,ₜ")
    add_para(doc, "uₙ,ₜ为高惩罚失供变量，用于识别边界不可行而不是作为常规供热手段。储热状态满足：")
    add_equation(doc, "Eₛ,ₜ = (1−μₛ)Eₛ,ₜ₋₁ + ηᶜʰₛ qᶜʰₛ,ₜΔt − qᵈʲˢₛ,ₜΔt/ηᵈʲˢₛ")
    add_para(doc, "循环边界按“末时段状态是首时段之前的状态”重建，避免把首小时充放热误报为跨年漂移。")
    add_heading(doc, "3.2 设备、管网与可靠性约束", 2)
    add_para(doc, "每个设备同时受容量、最小负荷、爬坡、启停和环境性能曲线约束。热泵COP按温度分段线性化，管网以定向能流与损失系数表达，可逆连络线通过方向状态防止同时双向流动。N−1容量门将最大单机故障后的剩余可靠容量与峰值负荷比较，与逐时失供约束共同组成“规划容量+运行可行”双重保障。")
    add_heading(doc, "3.3 多目标统一台账", 2)
    add_para(doc, "目标函数不把经济、环境和可靠性拆成相互独立的“多个算法”，而是统一到同一年度净成本台账：")
    add_equation(doc, "min Z = Cⁱⁿᵛ(x)+Eₛ[Qₛ(x,uₛ)]+λ·CVaRα(Qₛ)+ωᶜCᶜᵃʳᵇᵒⁿ+ωᵘCᵘⁿˢᵉʳᵛᵉᵈ")
    add_para(doc, "其中Qₛ已完整计入燃料、运维、碳、启停、泵耗、本地可再生电力抵扣和舒适度成本。CVaR以辅助变量η和zₛ建图：")
    add_equation(doc, "CVaRα(Q)=min{η+[1/(1−α)]Σₛ pₛzₛ},  zₛ≥Qₛ−η,  zₛ≥0")
    add_para(doc, "碳成本是决策项，碳排放实物量同时独立报告，从而避免把风险溢价或碳权价格误解为实际现金支出。")
    add_figure(doc, figures["heatmap"], "图3  源—荷匹配热力图（年度时序经日均聚合，保留季节结构）")

    add_heading(doc, "4 求解逻辑、收敛与效率", 1)
    add_heading(doc, "4.1 规划—回放—认证闭环", 2)
    add_para(doc, "第一阶段从多气象年中选取峰值、低温、平峰和资源波动代表时段，并用年度能量权重恢复季节结构；第二阶段冻结全部投资变量，在原始连续时序上重新优化运行；第三阶段用失供率、容量裕度、SOC边界、冷热可靠性、地源年度平衡和投资指纹生成机器可读证书。任意一道适用门失败，都不得将方案表述为“通过”。")
    add_heading(doc, "4.2 收敛与边界稳定性", 2)
    add_para(doc, "模型是稀疏混合整数线性规划。收敛证据同时报告solver status、MIP gap、对偶界、分支节点和约束残差，不以“运行结束”代替最优性说明。λ∈{0,0.02,0.05,0.10}的12个规划案例全部取得状态0解，60/60个“场景—权重—年度”容量审计通过，风险溢价与经验CVaR的最大一致性误差为5.59×10⁻⁹元。")
    add_figure(doc, figures["risk"], "图4  三场景成本—风险前沿（α=0.80）")
    add_heading(doc, "4.3 结构紧化与实测性能", 2)
    add_para(doc, "原启停模型将开机uₜ、启机sₜ和停机dₜ全部设为二元变量。新模型仅保留uₜ为二元，sₜ,dₜ为[0,1]连续量，并以三行凸包精确固定状态转移。这使全年单台设备的整数变量从26280个降至8760个，同时消除借用伪启停放松爬坡约束的可能。约束矩阵直接编译为CSR/CSC并缓存静态结构，重复求解无需重建边界与整数标记。")
    add_figure(doc, figures["performance"], "图5  固原极端年等投资方案求解性能对比")

    add_heading(doc, "5 仿真结果与多方案分析", 1)
    rows=[]
    for _, r in batch.iterrows():
        rows.append((r["city"], "8760", f"{int(r['variables']):,}", f"{r['solve_seconds']:.2f}", f"{r['clean_heat_share']*100:.1f}%", "0", f"{r['passed_gates']}/{r['gate_count']}"))
    add_table(doc, ["城市", "连续时数", "变量", "求解/s", "清洁热比例", "热失供率", "验收门"], rows, [2.2,2,2.8,2,2.4,2,2.1])
    add_heading(doc, "5.1 既有大型集中供热：乌鲁木齐", 2)
    add_para(doc, "该场景保留骨干管网与存量热电联产，将工业余热、余热源热泵和低温储热作为优先增量，燃气/电锅炉只承担峰值与故障备用。8760 h回放清洁热比例75.0%，热失供率0，可靠热容量裕度1100 kW。λ=0.05相对风险中性方案仅增加期望年成本124.29元，尾部运行成本降低8311.36元；λ继续增至0.10时边际收益明显放缓。")
    add_heading(doc, "5.2 分散式清洁转型：固原", 2)
    add_para(doc, "该场景以区域生物质、分布式空气源热泵、低温储热和微网可再生电力组成轻量级系统，用短距离区域管网完成“散烧—小集中—多能协同”过渡。8760 h回放清洁热比例91.8%、失供率0。λ=0.05是首个实质性投资转折：微网可再生容量增加20.46 kW，低温储热增加35.10 kWh，生物质锅炉减少3.13 kW，期望年成本增加156.93元，尾部成本降低5144.81元，代表层排放减少约2.33 t/a。")
    add_heading(doc, "5.3 夏热冬冷高品质新建：武汉", 2)
    add_para(doc, "该场景不复制北方大管网，而以分区空气源/地源热泵、太阳能热利用、PCM储热和可逆互联线建立冷热一体系统。8760 h回放热、冷失供率均为0，清洁热比例100%，地源取排热平衡比0.957，位于0.80–1.25验收区间。λ=0–0.10内投资指纹完全不变，说明当前方案主要由冷热容量、舒适度和地源平衡硬约束锁定。")
    add_heading(doc, "5.4 多气象年与极端年", 2)
    add_para(doc, "修正后的λ=0.05投资方案完成7个极端年连续回放：乌鲁木齐2024，固原2021/2024，武汉2021–2024，累计61392 h。所有回放状态为0，热/冷失供率均为0，7/7证书通过。这一结果证明投资并非只对单一年份或代表日有效，但仍不能替代未来气候情景、设备退化和价格扰动的样本外检验。")

    add_heading(doc, "6 经济性、碳排放与可行性评价", 1)
    add_para(doc, "年化经济台账将设备与管网投资通过资本回收因子折算，再加上燃料、电力、运维、启停、碳价和可靠性成本。对方案k，静现值NPV和投资回收期PBP按统一基准方案计算：")
    add_equation(doc, "NPVₖ=−Iₖ+Σᵧ₌₁ᴿ(ΔCᵒᵖ,ᵧ+ΔCᶜᵒ₂,ᵧ)/(1+r)ᵧ ;   PBPₖ=min{y:Σᵧ₌₁ʸΔCFᵧ≥Iₖ}")
    add_para(doc, "碳排放按项目边界内燃料直接排放和外购电力间接排放汇总，本地可再生电力替代电网时按同一排放因子抵扣。对外报告时应同时列出基准年、排放因子来源、边界和用能分母，未获得地方热力公司投资与结算数据前，本文的平准供热成本用于方案相对比较，不冒充财务审计结论。")
    add_table(doc, ["决策维度", "推荐判据", "风险边界"], [
        ("经济性", "用增量NPV、PBP、平准供热成本比较方案", "价格、折现率、利用小时和补贴必须做敏感性"),
        ("环境性", "报告实物碳排放与较基准减排量", "排放因子更新后必须重算"),
        ("可行性", "零失供+容量门+SOC边界+多年认证", "实测设备曲线、管网水力和当地电网接入待补"),
    ], [3, 7, 6])

    add_heading(doc, "7 代码质量与可移植性", 1)
    add_para(doc, "工程按“数据与配置—线性建模—场景适配—逐时回放—证书与报告”分层。主程序只组织流程，设备、网络、储能、风险和后检查均为可测试模块；城市差异通过YAML/CSV配置注入，避免在求解逻辑中写死地名。投资检查点、输入哈希、求解参数和结果证书使每次回放可追溯；87项自动化测试覆盖线性构件、时序边界、CVaR台账、启停语义和批量验收。")
    add_table(doc, ["层级", "职责", "移植接口"], [
        ("data/", "气象、价格、负荷、政策和城市参数", "CSV/JSON/YAML"),
        ("src/heatplan/", "统一建模、时序、风险、报告和验收", "Python API / CLI"),
        ("scripts/", "批量试验、绘图与工程交付", "可复现命令"),
        ("tests/", "单元、烟雾、回归与性能门禁", "pytest"),
        ("outputs/", "调度表、投资表、证书和图表", "专业规划平台可读"),
    ], [3.3, 7.2, 5.5])

    add_heading(doc, "8 结论与下一步", 1)
    add_para(doc, "本研究以一套统一的工程决策闭环覆盖三类供热转型场景，既保留城市间基建与资源差异，又使投资、逐时能流、用户舒适度和尾部风险在同一台账中比较。已完成的三场景8760 h回放与七个极端年认证证明模型具有连续时序、多目标收敛和多尺度耦合能力。对比不同风险权重后，λ=0.05是当前三场景的可解释统一折中，但不是脱离地方风险偏好的唯一最优值。")
    add_para(doc, "赛前最优先的三项工作是：第一，按4392/3984/2880 h运行2024–2025供暖季专项回放并替换“待回放”标识；第二，将公共原型负荷替换为热表、建筑和管网实测后重新校准；第三，对电价、余热衰减、低温COP、生物质价格和折现率执行样本外联合敏感性，报告方案失效概率而不只是单点最优值。")

    add_heading(doc, "参考文献", 1)
    refs = [
        "[1] Lund H, Werner S, Wiltshire R, et al. 4th Generation District Heating (4GDH). Energy, 2014, 68: 1–11.",
        "[2] Li H, Svendsen S. Energy and exergy analysis of low temperature district heating network. Energy, 2012, 45: 237–246.",
        "[3] Capone M, et al. Supply temperature reduction in district heating systems: methods and impacts. Energy, 2023.",
        "[4] Ceruti A, et al. Design optimization of renewable district heating with thermal energy storage. Energy Conversion and Management, 2025.",
        "[5] Chen X, et al. Multi-objective planning of integrated heating systems under renewable uncertainty. Applied Thermal Engineering, 2025.",
        "[6] Lyu Y, et al. Robust optimization for low-carbon integrated energy systems. Applied Energy, 2025.",
        "[7] Rockafellar R T, Uryasev S. Optimization of conditional value-at-risk. Journal of Risk, 2000, 2(3): 21–41.",
        "[8] 中华人民共和国国民经济和社会发展第十五个五年规划纲要, 2026.",
        "[9] 国务院. 美丽中国建设“十五五”规划（国发〔2026〕20号）, 2026.",
        "[10] “十五五”碳达峰行动方案, 2026.",
    ]
    for ref in refs:
        p=doc.add_paragraph(ref); p.paragraph_format.left_indent=Cm(.5); p.paragraph_format.hanging_indent=Cm(.5); p.paragraph_format.space_after=Pt(2)

    doc.add_page_break()
    add_heading(doc, "附录A 模型、数据与验收说明", 1)
    add_para(doc, "附录不重复正文结论，用于评审复核。完整变量与约束定义见代码src/heatplan/model.py；每个回放目录保留summary.json、dispatch.csv、investments.csv和annual_replay_acceptance_certificate.json。证书门包括：完整时轴、规划/回放求解状态、投资冻结、热/电容量裕度、失供率、储热边界、清洁热比例，以及武汉专属的供冷和地源平衡。")
    add_heading(doc, "附录B 计算结果索引", 1)
    add_table(doc, ["证据", "路径", "用途"], [
        ("三场景8760 h总表", "outputs/v11_5_three_scenario_annual_replay/annual_replay_batch_summary.csv", "模型规模、耗时、失供与验收"),
        ("CVaR扫描", "outputs/v11_10_cvar_weight_scan/risk_weight_scan_summary.csv", "多目标前沿与投资稳定性"),
        ("极端年回放", "outputs/v11_10_corrected_cvar_formal_replay/", "7年度证书"),
        ("性能验收", "outputs/v11_11_performance/", "整数变量、矩阵与墙钟时间"),
        ("供暖季范围", "data/heating_season_scopes_2024_2025.json", "4392/3984/2880 h机器可读边界"),
    ], [3.5, 8.5, 4])

    doc.core_properties.title = "暖城向新·三类供热转型通用优化模型竞赛研究报告"
    doc.core_properties.subject = "负荷预测、多尺度源荷匹配、多目标优化、逐时回放、经济与碳评价"
    doc.core_properties.author = "参赛团队"
    doc.save(DOCX)


def export_summary_csv(batch, risk, multi):
    out = batch[["profile","scenario","city","variables","integer_variables","constraints","solve_seconds","unserved_rate","clean_heat_share","cooling_unserved_rate","certified"]].copy()
    out["evidence_scope"] = "2023 continuous 8760h fixed-investment replay"
    out["heating_season_hours_2024_2025"] = out["city"].map({"乌鲁木齐":4392,"固原":3984,"武汉":2880})
    out["heating_season_replay_status"] = "pending; not claimed as passed"
    out.to_csv(CSV, index=False, encoding="utf-8-sig")


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    batch, risk, multi, perf = load_data()
    figures = make_figures(batch, risk, multi, perf)
    export_summary_csv(batch, risk, multi)
    build_doc(batch, risk, multi, figures)
    print(DOCX)
    print(CSV)


if __name__ == "__main__":
    main()
