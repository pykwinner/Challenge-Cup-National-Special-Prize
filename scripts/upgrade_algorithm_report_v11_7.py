from __future__ import annotations

from pathlib import Path

from docx import Document

from upgrade_algorithm_report_v11_1 import add_paragraph_before
from upgrade_algorithm_report_v11_3 import align_table_indents
from upgrade_algorithm_report_v11_4 import add_table_before


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "为什么数学建模能够支撑供热转型决策_算法增强版V8.docx"
OUTPUT = ROOT / "为什么数学建模能够支撑供热转型决策_算法增强版V9.docx"


def refresh_summary(doc) -> None:
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            if "为什么可信：V11.6" in paragraph.text and "V11.6" in run.text:
                run.text = run.text.replace("V11.6", "V11.7")
            if (
                "全工程73项" in paragraph.text
                and not paragraph.text.strip().startswith("V11.6全工程")
                and "73" in run.text
            ):
                run.text = run.text.replace("73", "79")


def add_v11_7_section(doc) -> None:
    anchor = next(
        paragraph
        for paragraph in doc.paragraphs
        if paragraph.text.strip().startswith("3  为什么该模型具有工程物理可信性")
    )

    add_paragraph_before(
        doc,
        anchor,
        "2.7.14  V11.7：官方时政核验与场景化政策证据链",
        "Heading 3",
    )
    add_paragraph_before(
        doc,
        anchor,
        "V11.7以2026年3月13日发布的国家“十五五”规划纲要和国发〔2026〕22号《“十五五”碳达峰行动方案》为正式依据，将政策内容从报告背景升级为可校验输入。政策画像记录官网、发布机关、文件号、发布日期、核验日期、适用场景、模型字段、证据要求和口径限制；标记为“官方已核验”的来源必须使用gov.cn政府官网域名。国家宏观目标默认只作战略参考，只有地方任务已分解、2025基线经审计且分母一致的项目KPI，才允许输出合规结论。",
        "Normal",
    )
    add_table_before(
        doc,
        anchor,
        ["正式文件", "核验信息", "进入模型的主要机制"],
        [
            (
                "国民经济和社会发展第十五个五年规划纲要",
                "2026-03-13\n中国政府网",
                "能源安全、有序替代、余热和非化石供热、供热计量、城乡配网、城市更新与老旧管网生命线",
            ),
            (
                "“十五五”碳达峰行动方案",
                "国发〔2026〕22号\n2026-07-05",
                "清洁热力源荷匹配、热电解耦、储热蓄冷、分时价格、需求响应、建筑直接碳排和小燃煤锅炉淘汰",
            ),
        ],
        [2450, 1800, 5110],
    )
    add_paragraph_before(
        doc,
        anchor,
        "对政策条款p和供热场景s，模型先计算适用指示量a(p,s)。若a=0，输出“不适用”而不是“不合规”；若a=1，则按官方来源、模型实现、地方参数、审计基线和资产清单依次检查。项目合规判定必须同时满足official、local target、baseline、denominator和implemented五个条件。该门槛防止把全国GDP碳强度17%、非化石能源消费25%、建筑直接碳排3%或全国需求响应5%机械改写成单个供热项目阈值。",
        "Normal",
    )
    add_table_before(
        doc,
        anchor,
        ["场景", "模型就绪", "就绪但证据待补", "关键证据缺口", "战略参考", "不适用"],
        [
            ("S1 既有大型集中供热", "6", "2", "2", "4", "2"),
            ("S2 分散式清洁转型", "4", "2", "2", "4", "4"),
            ("S3 夏热冬冷高品质新建", "5", "2", "1", "4", "4"),
        ],
        [2700, 1200, 1700, 1500, 1160, 1100],
    )
    add_paragraph_before(
        doc,
        anchor,
        "16项政策条款完成三场景筛选后，S1和S2各形成4项行动任务，S3形成3项：三场景均需补齐2025建筑直接排放基线、服务面积和排放因子；S1/S2还需逐台燃煤锅炉蒸吨清单；三场景的供热计量、分时电价资格和需求响应补偿规则需要地方文件校准。三场景重新求解均成功，最坏场景失供率仍为0，说明政策元数据升级没有破坏原投资与运行逻辑。",
        "Normal",
    )

    add_paragraph_before(
        doc,
        anchor,
        "2.7.15  小燃煤锅炉逐台清单与强制退役硬约束",
        "Heading 3",
    )
    add_paragraph_before(
        doc,
        anchor,
        "国发〔2026〕22号要求全面淘汰10蒸吨/小时及以下燃煤锅炉。V11.7没有把“10蒸吨”粗略换算成城市总热功率，而是新增资产级输入boiler_capacity_t_per_hour。启用淘汰门后，每台存量化石锅炉都必须提供蒸吨容量；对Bᵢ≤10的锅炉，强制Rᵢ=Kᵢᵉˣⁱˢᵗ并使Kᵢʳᵉᵐᵃⁱⁿ=0。若最大可退役容量不足以覆盖存量容量，系统在构模前直接拒绝输入。",
        "Normal",
    )
    add_paragraph_before(
        doc,
        anchor,
        "该约束解决三类常见政策漂移：只在文字中宣称淘汰但模型仍调度小锅炉；用系统总供热规模替代逐台锅炉容量；为了降低成本而保留政策要求退出的设备。门槛默认关闭，因为当前公开示范数据尚无完整逐台资产台账；正式竞赛或城市项目应先补齐锅炉容量、环保、能耗、安全、年限和地方实施计划，再显式开启。",
        "Normal",
    )
    add_paragraph_before(
        doc,
        anchor,
        "V11.7全工程79项自动化测试通过。新增测试覆盖官方来源域名、日期与文件元数据、场景适用性、项目KPI证据门、政策就绪度汇总、行动清单、逐台锅炉清单缺失拒绝和小锅炉全容量强制退役。由此，模型不仅回答“哪种技术成本最低”，还能够回答“该政策是否适用、缺什么证据、何时允许形成项目硬约束”。",
        "Normal",
    )


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    doc = Document(SOURCE)
    refresh_summary(doc)
    add_v11_7_section(doc)
    align_table_indents(doc)
    doc.core_properties.title = "为什么数学建模能够支撑供热转型决策——算法增强版V9"
    doc.core_properties.subject = "V11.7官方时政核验、场景化政策证据链与小燃煤锅炉强制退役"
    doc.core_properties.version = "11.7"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
